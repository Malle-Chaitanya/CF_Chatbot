# -*- coding: utf-8 -*-
"""
Export CloudFuze Users Script

Fetches all users from the CloudFuze domain using Microsoft Graph API
and exports their names, emails, and job titles to CSV and JSON files.
"""

import os
import sys
import requests
import csv
import json
from typing import List, Dict
from datetime import datetime
from dotenv import load_dotenv

# Load environment variables from .env file
load_dotenv()

# Set UTF-8 encoding for Windows console
if sys.platform == 'win32':
    import io
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')

# Add parent directory to path to import app modules
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app.sharepoint_auth import sharepoint_auth


def get_all_cloudfuze_users() -> List[Dict]:
    """
    Fetch all CloudFuze domain users with their details using Microsoft Graph API.
    
    Returns:
        List of user dictionaries containing id, displayName, userPrincipalName, mail, and jobTitle
    """
    print("[*] Connecting to Microsoft Graph API...")
    
    try:
        headers = sharepoint_auth.get_headers()
        # Add ConsistencyLevel header for advanced queries
        headers["ConsistencyLevel"] = "eventual"
        
        base_url = "https://graph.microsoft.com/v1.0/users"
        
        # Query for all users in CloudFuze domain
        params = {
            "$filter": "endswith(userPrincipalName,'@cloudfuze.com')",
            "$select": "id,displayName,userPrincipalName,mail,jobTitle",
            "$top": 999,  # Graph API max per request
            "$count": "true"
        }
        
        all_users = []
        page_count = 0
        url = base_url
        
        while url:
            page_count += 1
            print(f"[*] Fetching page {page_count}...")
            
            response = requests.get(url, headers=headers, params=params if page_count == 1 else None, timeout=30)
            response.raise_for_status()
            
            data = response.json()
            users = data.get("value", [])
            all_users.extend(users)
            
            print(f"    -> Got {len(users)} users (Total: {len(all_users)})")
            
            # Handle pagination
            url = data.get("@odata.nextLink")
            # Clear params for subsequent requests (they're included in nextLink)
            params = None
        
        print(f"[OK] Total users fetched: {len(all_users)}")
        return all_users
        
    except requests.exceptions.HTTPError as e:
        print(f"[ERROR] Failed to fetch users from Graph API: {e}")
        if hasattr(e, 'response') and e.response is not None:
            response_text = e.response.text
            print(f"[ERROR] Response: {response_text}")
            
            # Check for permission errors
            if "403" in str(e) or "Insufficient privileges" in response_text:
                print("\n" + "=" * 70)
                print("PERMISSION ERROR - SETUP REQUIRED")
                print("=" * 70)
                print("You need to add 'User.Read.All' permission to your Azure AD app:")
                print("\n1. Go to Azure Portal: https://portal.azure.com")
                print("2. Navigate to: Azure Active Directory > App registrations")
                print("3. Find your app (check MICROSOFT_CLIENT_ID)")
                print("4. Click 'API permissions'")
                print("5. Click 'Add a permission'")
                print("6. Select 'Microsoft Graph'")
                print("7. Choose 'Application permissions'")
                print("8. Search for and select 'User.Read.All'")
                print("9. Click 'Add permissions'")
                print("10. Click 'Grant admin consent for [your tenant]'")
                print("11. Confirm by clicking 'Yes'")
                print("\nAfter adding the permission, wait 1-2 minutes for it to propagate.")
                print("Then run this script again.")
                print("=" * 70)
        raise
    except requests.exceptions.RequestException as e:
        print(f"[ERROR] Failed to fetch users from Graph API: {e}")
        raise
    except Exception as e:
        print(f"[ERROR] Unexpected error: {e}")
        raise


def save_to_csv(users: List[Dict], filename: str) -> int:
    """
    Save users to CSV file.
    
    Args:
        users: List of user dictionaries
        filename: Output CSV file path
        
    Returns:
        Number of users saved
    """
    print(f"\n[*] Saving to CSV: {filename}")
    
    try:
        with open(filename, 'w', newline='', encoding='utf-8') as f:
            writer = csv.writer(f)
            writer.writerow(['Name', 'Email', 'Position', 'User ID'])
            
            for user in users:
                writer.writerow([
                    user.get('displayName', ''),
                    user.get('mail') or user.get('userPrincipalName', ''),
                    user.get('jobTitle', '') or 'N/A',
                    user.get('id', '')
                ])
        
        print(f"[OK] Saved {len(users)} users to {filename}")
        return len(users)
        
    except Exception as e:
        print(f"[ERROR] Failed to save CSV: {e}")
        raise


def save_to_json(users: List[Dict], filename: str) -> int:
    """
    Save users to JSON file.
    
    Args:
        users: List of user dictionaries
        filename: Output JSON file path
        
    Returns:
        Number of users saved
    """
    print(f"[*] Saving to JSON: {filename}")
    
    try:
        output_data = {
            "export_date": datetime.now().isoformat(),
            "total_users": len(users),
            "users": users
        }
        
        with open(filename, 'w', encoding='utf-8') as f:
            json.dump(output_data, f, indent=2, ensure_ascii=False)
        
        print(f"[OK] Saved {len(users)} users to {filename}")
        return len(users)
        
    except Exception as e:
        print(f"[ERROR] Failed to save JSON: {e}")
        raise


def save_to_excel(users: List[Dict], filename: str) -> int:
    """
    Save users to Excel file (if openpyxl is available).
    
    Args:
        users: List of user dictionaries
        filename: Output Excel file path
        
    Returns:
        Number of users saved
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        
        print(f"[*] Saving to Excel: {filename}")
        
        # Create a new workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "CloudFuze Users"
        
        # Add headers with styling
        headers = ['Name', 'Email', 'Position', 'User ID']
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Add user data
        for row_num, user in enumerate(users, 2):
            ws.cell(row=row_num, column=1).value = user.get('displayName', '')
            ws.cell(row=row_num, column=2).value = user.get('mail') or user.get('userPrincipalName', '')
            ws.cell(row=row_num, column=3).value = user.get('jobTitle', '') or 'N/A'
            ws.cell(row=row_num, column=4).value = user.get('id', '')
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 35
        ws.column_dimensions['C'].width = 25
        ws.column_dimensions['D'].width = 40
        
        # Save workbook
        wb.save(filename)
        print(f"[OK] Saved {len(users)} users to {filename}")
        return len(users)
        
    except ImportError:
        print("[WARNING] openpyxl not available, skipping Excel export")
        return 0
    except Exception as e:
        print(f"[ERROR] Failed to save Excel: {e}")
        return 0


def print_summary(users: List[Dict]):
    """Print a summary of exported users."""
    print("\n" + "=" * 70)
    print("CLOUDFUZE USERS EXPORT SUMMARY")
    print("=" * 70)
    print(f"Total Users: {len(users)}")
    
    # Stats
    with_position = sum(1 for u in users if u.get('jobTitle'))
    without_position = len(users) - with_position
    
    print(f"  -> With Position: {with_position}")
    print(f"  -> Without Position: {without_position}")
    
    # Sample users
    if users:
        print("\nSample Users (first 5):")
        for i, user in enumerate(users[:5], 1):
            email = user.get('mail') or user.get('userPrincipalName', 'N/A')
            position = user.get('jobTitle', 'N/A')
            print(f"  {i}. {user.get('displayName', 'N/A')} <{email}> - {position}")
    
    print("=" * 70)


def main():
    """Main function to orchestrate the export."""
    print("\n" + "=" * 70)
    print("CLOUDFUZE USERS EXPORTER")
    print("=" * 70)
    
    try:
        # Fetch users
        print("\n[STEP 1] Fetching CloudFuze users from Microsoft Graph API...")
        users = get_all_cloudfuze_users()
        
        if not users:
            print("[WARNING] No users found!")
            return
        
        # Determine output directory
        output_dir = os.path.dirname(os.path.abspath(__file__))
        
        # Save to different formats
        print("\n[STEP 2] Exporting to files...")
        
        csv_file = os.path.join(output_dir, "cloudfuze_users.csv")
        json_file = os.path.join(output_dir, "cloudfuze_users.json")
        excel_file = os.path.join(output_dir, "cloudfuze_users.xlsx")
        
        save_to_csv(users, csv_file)
        save_to_json(users, json_file)
        save_to_excel(users, excel_file)
        
        # Print summary
        print_summary(users)
        
        print("\n[OK] Export complete!")
        print(f"\nOutput files created in: {output_dir}")
        print(f"  - {os.path.basename(csv_file)}")
        print(f"  - {os.path.basename(json_file)}")
        print(f"  - {os.path.basename(excel_file)}")
        
    except Exception as e:
        print(f"\n[ERROR] Export failed: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)


if __name__ == "__main__":
    main()




